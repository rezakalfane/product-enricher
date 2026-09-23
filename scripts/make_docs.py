"""
Generate realistic supplier documents (.docx, .pdf, .csv) that reference real products
from the channel, plus a few invented ones so the "create product" path gets exercised.

Run:  ./.venv-fixtures/bin/python scripts/make_docs.py
"""
import csv, json, random, pathlib

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, HRFlowable)

random.seed(20260922)
OUT = pathlib.Path("fixtures/docs"); OUT.mkdir(parents=True, exist_ok=True)
CATALOG = json.loads(pathlib.Path("fixtures/catalog-sample.json").read_text())

BRAND_BLUE = colors.HexColor("#3C64F4")
INK = colors.HexColor("#121118")

# --- Enrichment values the catalog does not currently carry -------------------
FLAVOUR_NOTES = ["Rich espresso over steamed milk with a cocoa finish",
                 "Crisp orchard apple with a menthol exhale",
                 "Blackcurrant sorbet, iced",
                 "Sour blue raspberry with a sherbet lift",
                 "Dark berries, aniseed and liquorice",
                 "Sweet banana cream over ice"]
VG_PG = ["70VG / 30PG", "50VG / 50PG", "60VG / 40PG", "80VG / 20PG"]
BOTTLES = ["100ml chubby gorilla unicorn bottle", "200ml short fill bottle",
           "10ml TPD-compliant bottle with child lock"]
ORIGINS = ["United Kingdom", "Malaysia", "United States"]
STEEP = ["No steeping required", "Steep 3-5 days for best results", "Shake well before use"]

def pick(n, predicate=lambda p: True):
    pool = [p for p in CATALOG if predicate(p)]
    return random.sample(pool, min(n, len(pool)))

def enrichment(p):
    """A plausible set of attributes a supplier would send for this product."""
    return {
        "Flavour profile": random.choice(FLAVOUR_NOTES),
        "VG/PG ratio": random.choice(VG_PG),
        "Bottle format": random.choice(BOTTLES),
        "Country of origin": random.choice(ORIGINS),
        "Steeping": random.choice(STEEP),
        "Nicotine strength": random.choice(["0mg", "3mg", "6mg", "10mg", "20mg"]),
        "Shelf life": f"{random.choice([18, 24, 36])} months from manufacture",
        "Case quantity": str(random.choice([6, 10, 12, 24])),
    }

# --- Invented products, for the "create product" path -------------------------
NEW_PRODUCTS = [
    dict(name="Nordsund Arctic Cloudberry 0mg 100ml Shortfill E-Liquid",
         sku="NRDSND-ACB-100", brand="Nordsund", price=12.99, weight=0.14,
         mpn="NS-ACB-100-0",
         attrs={"Flavour profile": "Wild cloudberry with a soft vanilla base",
                "VG/PG ratio": "70VG / 30PG", "Bottle format": "100ml short fill",
                "Nicotine strength": "0mg", "Country of origin": "Norway",
                "Shelf life": "24 months from manufacture"}),
    dict(name="Halden Vapour Works Rhubarb Custard Nic Salt 10ml 20mg",
         sku="HVW-RHC-10-20", brand="Halden Vapour Works", price=3.99, weight=0.03,
         mpn="HVW-RHC-1020",
         attrs={"Flavour profile": "Stewed rhubarb folded into thick custard",
                "VG/PG ratio": "50VG / 50PG", "Bottle format": "10ml TPD bottle",
                "Nicotine strength": "20mg nicotine salt",
                "Country of origin": "United Kingdom", "Case quantity": "10"}),
    dict(name="Kestrel Mod Co. Falcon X Pod Kit",
         sku="KMC-FALCONX", brand="Kestrel Mod Co.", price=34.95, weight=0.21,
         mpn="KMC-FX-2026",
         attrs={"Battery capacity": "1500mAh integrated",
                "Output range": "5W - 30W adjustable",
                "Pod capacity": "3ml refillable",
                "Charging": "USB-C, 1.5A fast charge",
                "Coil compatibility": "Kestrel FX series, 0.4ohm and 0.8ohm",
                "Country of origin": "China"}),
]

# ============================== DOCX ==========================================
def style_doc(doc, title, subtitle):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(10.5)
    h = doc.add_paragraph()
    r = h.add_run(title); r.bold = True; r.font.size = Pt(19)
    r.font.color.rgb = RGBColor(0x12, 0x11, 0x18)
    s = doc.add_paragraph()
    sr = s.add_run(subtitle); sr.font.size = Pt(9.5)
    sr.font.color.rgb = RGBColor(0x6C, 0x6A, 0x7A)
    return doc

def add_kv_table(doc, rows):
    t = doc.add_table(rows=0, cols=2); t.style = "Light Grid Accent 1"
    t.autofit = False
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].width = Inches(2.0); cells[1].width = Inches(4.2)
        run = cells[0].paragraphs[0].add_run(k); run.bold = True
        cells[1].paragraphs[0].add_run(str(v))
    doc.add_paragraph()

def doc_spec_sheet(products, filename, supplier):
    doc = Document()
    style_doc(doc, f"{supplier} — Product Specification Sheet",
              "Confidential · For retail partner catalog use · Revision 2026.09")
    doc.add_paragraph(
        "The attributes below supersede any previously issued specification. "
        "Please reflect these values on all listings within 14 days."
    )
    for p in products:
        doc.add_paragraph()
        hp = doc.add_paragraph(); hr = hp.add_run(p["name"])
        hr.bold = True; hr.font.size = Pt(13)
        hr.font.color.rgb = RGBColor(0x3C, 0x64, 0xF4)
        rows = [("Supplier SKU", p["sku"]), ("Trade price", f"GBP {p['price']:.2f}")]
        rows += list(enrichment(p).items())
        add_kv_table(doc, rows)
    doc.save(OUT / filename)
    print("  ✓", filename)

def doc_new_launch(filename):
    doc = Document()
    style_doc(doc, "New Product Launch Pack — Autumn 2026",
              "Nordsund Distribution Ltd · Trade confidential")
    doc.add_paragraph(
        "Three new lines are available to order from 1 October 2026. None of these "
        "have previously been listed with you. Full specifications follow."
    )
    for p in NEW_PRODUCTS:
        doc.add_paragraph()
        hp = doc.add_paragraph(); hr = hp.add_run(p["name"])
        hr.bold = True; hr.font.size = Pt(13)
        hr.font.color.rgb = RGBColor(0x00, 0xA3, 0xA3)
        rows = [("Brand", p["brand"]), ("SKU", p["sku"]), ("MPN", p["mpn"]),
                ("RRP", f"GBP {p['price']:.2f}"), ("Unit weight", f"{p['weight']} kg")]
        rows += list(p["attrs"].items())
        add_kv_table(doc, rows)
    doc.save(OUT / filename)
    print("  ✓", filename)

def doc_supplier_email(product, filename):
    doc = Document()
    style_doc(doc, "FW: Updated product data — please action",
              "From: merchandising@bellwether-wholesale.co.uk · 22 September 2026")
    doc.add_paragraph("Hi,")
    doc.add_paragraph(
        f"Following our call, here is the corrected data for {product['name']} "
        f"(our ref {product['sku']}). The listing currently has no flavour "
        "description and the bottle format is missing, which is hurting conversion."
    )
    for k, v in enrichment(product).items():
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")
    doc.add_paragraph(
        f"Trade price stays at GBP {product['price']:.2f}. Let me know once it is live."
    )
    doc.add_paragraph("Best,\nSam Oduya\nAccount Manager, Bellwether Wholesale")
    doc.save(OUT / filename)
    print("  ✓", filename)

# ============================== PDF ===========================================
styles = getSampleStyleSheet()
H1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=18, textColor=INK,
                    spaceAfter=2, fontName="Helvetica-Bold")
SUB = ParagraphStyle("SUB", parent=styles["Normal"], fontSize=8.5,
                     textColor=colors.HexColor("#6C6A7A"), spaceAfter=10)
H2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=12,
                    textColor=BRAND_BLUE, spaceBefore=12, spaceAfter=4)
BODY = ParagraphStyle("BODY", parent=styles["Normal"], fontSize=9.5, leading=13.5)

TBL = TableStyle([
    ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
    ("FONTSIZE", (0, 0), (-1, -1), 9),
    ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#34313F")),
    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E3E5EC")),
    ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#F6F7F9")]),
    ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
])

def build_pdf(filename, title, subtitle, blocks):
    doc = SimpleDocTemplate(str(OUT / filename), pagesize=A4,
                            leftMargin=18*mm, rightMargin=18*mm,
                            topMargin=16*mm, bottomMargin=16*mm, title=title)
    flow = [Paragraph(title, H1), Paragraph(subtitle, SUB),
            HRFlowable(width="100%", color=BRAND_BLUE, thickness=1.6), Spacer(1, 8)]
    flow += blocks
    doc.build(flow)
    print("  ✓", filename)

def pdf_spec_sheet(products, filename, supplier):
    blocks = [Paragraph(
        "Specification data issued for catalog enrichment. Values marked here are "
        "authoritative and replace anything previously supplied.", BODY)]
    for p in products:
        rows = [["Supplier SKU", p["sku"]], ["Trade price", f"GBP {p['price']:.2f}"]]
        rows += [[k, v] for k, v in enrichment(p).items()]
        blocks += [Paragraph(p["name"], H2),
                   Table(rows, colWidths=[45*mm, 120*mm], style=TBL)]
    build_pdf(filename, f"{supplier} — Specification Sheet",
              "Trade confidential · Revision 2026.09 · Page 1 of 1", blocks)

def pdf_price_list(products, filename):
    head = ["SKU", "Product", "Nic strength", "Case qty", "Trade GBP", "RRP GBP"]
    rows = [head]
    for p in products:
        e = enrichment(p)
        rows.append([p["sku"], p["name"][:52], e["Nicotine strength"],
                     e["Case quantity"], f"{p['price']:.2f}", f"{p['price'] * 1.9:.2f}"])
    style = TableStyle(TBL.getCommands() + [
        ("BACKGROUND", (0, 0), (-1, 0), BRAND_BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (1, 1), (1, -1), "Helvetica"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F7F9")]),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
    ])
    blocks = [Paragraph(
        "Prices effective 1 October 2026 and held until 31 December 2026. "
        "Case quantities are minimum order multiples.", BODY), Spacer(1, 8),
        Table(rows, colWidths=[28*mm, 62*mm, 24*mm, 18*mm, 20*mm, 20*mm], style=style)]
    build_pdf(filename, "Q4 2026 Trade Price List",
              "Bellwether Wholesale Ltd · All prices exclude VAT", blocks)

def pdf_new_products(filename):
    blocks = [Paragraph(
        "The following lines are new to the range and are not currently listed. "
        "Please create them ahead of the 1 October on-sale date.", BODY)]
    for p in NEW_PRODUCTS:
        rows = [["Brand", p["brand"]], ["SKU", p["sku"]], ["MPN", p["mpn"]],
                ["RRP", f"GBP {p['price']:.2f}"], ["Unit weight", f"{p['weight']} kg"]]
        rows += [[k, v] for k, v in p["attrs"].items()]
        blocks += [Paragraph(p["name"], H2),
                   Table(rows, colWidths=[45*mm, 120*mm], style=TBL)]
    build_pdf(filename, "New Line Notification — Autumn 2026",
              "Nordsund Distribution Ltd · Trade confidential", blocks)

# ============================== CSV ===========================================
def csv_enrichment(products, filename):
    with open(OUT / filename, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["supplier_sku", "product_name", "flavour_profile", "vg_pg_ratio",
                    "bottle_format", "nicotine_strength", "country_of_origin",
                    "shelf_life", "case_qty", "trade_price_gbp"])
        for p in products:
            e = enrichment(p)
            w.writerow([p["sku"], p["name"], e["Flavour profile"], e["VG/PG ratio"],
                        e["Bottle format"], e["Nicotine strength"],
                        e["Country of origin"], e["Shelf life"], e["Case quantity"],
                        f"{p['price']:.2f}"])
    print("  ✓", filename)

# ============================== run ===========================================
print("Generating fixture documents in fixtures/docs/")
a, b, c = pick(4), pick(3), pick(12)
doc_spec_sheet(a, "bellwether-spec-sheet-shortfills.docx", "Bellwether Wholesale")
doc_spec_sheet(b, "vaporline-spec-sheet-kits.docx", "Vaporline Trade")
doc_new_launch("nordsund-new-line-pack.docx")
doc_supplier_email(pick(1)[0], "supplier-email-data-correction.docx")
pdf_spec_sheet(pick(4), "vaporline-specification-sheet.pdf", "Vaporline Trade")
pdf_price_list(c, "q4-2026-trade-price-list.pdf")
pdf_new_products("nordsund-new-line-notification.pdf")
csv_enrichment(pick(20), "supplier-enrichment-feed.csv")
print(f"\nDone → {OUT}/")
